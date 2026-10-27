import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

SAIDA = r"C:/LIA - UEFS/motores-video/saida/receitas-bistro/Receitas-do-Bistro-controle.xlsx"
F = "Arial"
VINHO, CREME, MARROM = "8D0C00", "FFFDEA", "3D2F29"


def fonte(size=10, **k):
    return Font(name=F, size=size, **k)


INPUT = fonte(color="0000FF")
LINK = fonte(color="008000")
FORM = fonte(color="000000")
AMARELO = PatternFill("solid", fgColor="FFF2A8")
CABEC = PatternFill("solid", fgColor=VINHO)
TOPO = PatternFill("solid", fgColor=MARROM)
SUAVE = PatternFill("solid", fgColor="F6EFD9")
fino = Side(style="thin", color="D9CFB8")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
RS = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
PCT = '0.0%;-0.0%;"-"'
QTD = '#,##0;-#,##0;"-"'
X = '0.0"x";-0.0"x";"-"'
DATA = "DD/MM/YYYY"

wb = Workbook()


def titulo(ws, texto, sub, larg):
    ws["A1"] = texto
    ws["A1"].font = fonte(14, bold=True, color=CREME)
    ws["A2"] = sub
    ws["A2"].font = fonte(9, italic=True, color=MARROM)
    for c in range(1, larg + 1):
        ws.cell(1, c).fill = CABEC
    ws.row_dimensions[1].height = 24
    ws.sheet_view.showGridLines = False


def cabecalho(ws, linha, nomes, larguras, congela=True):
    for i, (n, w) in enumerate(zip(nomes, larguras), 1):
        c = ws.cell(linha, i, n)
        c.font = fonte(bold=True, color=CREME)
        c.fill = TOPO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[linha].height = 32
    if congela:
        ws.freeze_panes = ws.cell(linha + 1, 1)


def put(ws, ref, v, font=FORM, fmt=None, fill=None, borda=True):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if borda:
        c.border = BORDA
    return c


# ============================ Leia-me ============================
ws = wb.active
ws.title = "Leia-me"
titulo(ws, "Receitas do Bistrô · controle financeiro e de estoque",
       "Bistrô 153 · Serrinha (BA) · planilha criada em " + dt.date.today().strftime("%d/%m/%Y"), 2)
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 104
linhas = [
    ("COMO USAR", None),
    ("1. Vendas", 'Uma linha por venda: data, código, quantidade, desconto, pagamento, canal e "Refil?". Brinde: canal "Brinde" (sai do estoque e entra no custo, mas não na receita).'),
    ("2. Movimentos", "Toda produção entra aqui (Tipo = Produção). Perdas, quebras e potes usados no restaurante: Tipo = Perda ou Ajuste, com quantidade negativa."),
    ("3. Estoque e Resumo", "Só leitura: saldo de cada produto, aviso de reposição e os números de cada mês (bruto, líquido, lucro, resultado)."),
    ("4. Produtos", "Preencha o custo do conteúdo (ingredientes) de cada produto, o preço de venda e o desconto do refil. Custo, lucro e margem são calculados."),
    ("5. Kits", "Ajuste quantos de cada produto vão em cada kit e o preço. O custo do kit soma os produtos, a caixa, a fita e a tag."),
    ("6. Insumos", "Atualize o preço pago e o tamanho da compra de cada insumo. O custo por unidade é calculado e alimenta os produtos."),
    ("7. Despesas", "Gastos que NÃO são custo de produto: divulgação, impressão de cartões, alvará, entregas, equipamentos."),
    ("", None),
    ("REFIL", None),
    ("Como funciona", 'O cliente devolve o pote ou a garrafa limpa e leva cheio de novo com desconto. Na venda, marque "Sim" na coluna Refil? O desconto de refil do produto (aba Produtos) sai sozinho do bruto.'),
    ("Custo do refil", "Na venda de refil, o custo NÃO inclui a embalagem (o pote volta), mas soma a higienização (aba Parâmetros). O rótulo continua no custo: troque o rótulo se precisar."),
    ("Por que vale", "O desconto é menor que o preço do pote, então o Bistrô ganha por venda de refil e o cliente volta. Conta no Resumo quantos refis cada produto teve."),
    ("", None),
    ("CORES", None),
    ("Azul", "Valor digitado por você (pode mudar)."),
    ("Fundo amarelo", "ESTIMATIVA, ainda não confirmada: revise com os preços reais de vocês."),
    ("Verde", "Valor puxado de outra aba."),
    ("Preto", "Fórmula: não precisa mexer."),
    ("", None),
    ("CONCEITOS", None),
    ("Bruto", "Quanto o cliente pagou: quantidade × preço − desconto − desconto de refil."),
    ("Líquido", "Bruto menos a taxa da forma de pagamento (cartão, link). Pix e dinheiro: sem taxa."),
    ("Lucro da venda", "Líquido menos o custo dos produtos vendidos (conteúdo + embalagem + rótulo + gás)."),
    ("Resultado do mês", "Lucro das vendas menos as despesas do mês. Compras de insumos NÃO entram em Despesas: já estão no custo de cada produto."),
    ("Kits e estoque", "Ao vender um kit, o estoque de cada produto do kit baixa sozinho (pela composição da aba Kits). Se o cliente trocar um sabor, registre em Movimentos (Ajuste: −1 do sabor que saiu do kit padrão, +1 do que entrou)."),
    ("", None),
    ("DE ONDE VÊM OS NÚMEROS", None),
    ("Informados por você", "Cachaça-base R$ 7,00/L · pote de 40 ml R$ 2,00 · garrafa da cachaça de 1 L."),
    ("Pesquisados", "Pote de 240 ml com tampa metálica e lacre: R$ 3,92 a R$ 5,33 por unidade conforme a quantidade (out/2026). Usado R$ 4,40."),
    ("Estimativas", "Demais insumos, custo do conteúdo dos potes, taxas de cartão, descontos de refil e preços sugeridos. Referência de mercado (out/2026): geleia artesanal 40–50 g R$ 8–18; 240 g R$ 22–40; cachaça com mel e limão 700–900 ml R$ 50–65."),
    ("Atenção", "Alho confit e pesto em óleo: conservar refrigerado. Cachaça com ingredientes (bebida mista) costuma exigir registro no MAPA para venda embalada; venda só para maiores de 18 anos."),
]
r = 4
cores = {"Azul": fonte(bold=True, color="0000FF"), "Verde": fonte(bold=True, color="008000"), "Preto": fonte(bold=True)}
for a, b in linhas:
    c = ws.cell(r, 1, a)
    c.font = fonte(bold=True, color=VINHO) if b is None else cores.get(a, fonte(bold=True, color=MARROM))
    if a == "Fundo amarelo":
        c.fill = AMARELO
    if b is not None:
        d = ws.cell(r, 2, b)
        d.font = fonte()
        d.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

# ============================ Parâmetros ============================
pm = wb.create_sheet("Parâmetros")
titulo(pm, "Parâmetros", "Taxas, listas e regras usadas nas outras abas. O catálogo (colunas H a M) é automático.", 13)
cabecalho(pm, 4, ["Forma de pagamento", "Taxa (%)", "", "Canal", "", "Regra"], [22, 11, 3, 18, 3, 40], congela=False)
pags = [("Pix", 0), ("Dinheiro", 0), ("Débito", 0.0199), ("Crédito à vista", 0.0349), ("Link de pagamento", 0.0499)]
for i, (n, t) in enumerate(pags):
    put(pm, f"A{5 + i}", n, INPUT)
    put(pm, f"B{5 + i}", t, INPUT, PCT, AMARELO if t else None)
pm["B7"].comment = Comment("Estimativa de taxa de maquininha. Troque pelas taxas do seu contrato.", "Planilha")
for i, n in enumerate(["Direct Instagram", "WhatsApp", "Balcão do Bistrô", "Encomenda", "Brinde"]):
    put(pm, f"D{5 + i}", n, INPUT)
put(pm, "F5", "Multiplicador alvo (preço ÷ custo)")
put(pm, "F6", 2.5, INPUT, X, AMARELO)
pm["F6"].comment = Comment("Regra de bolso para produto artesanal: preço = custo × 2,5 a 3.", "Planilha")
put(pm, "F8", "Brinde: mini 40 ml nas compras a partir de")
put(pm, "F9", 60, INPUT, RS)
put(pm, "F11", "Higienização por pote/garrafa de refil")
put(pm, "F12", 0.20, INPUT, RS, AMARELO)
pm["F12"].comment = Comment("Água, sabão, sanitizante e tempo para lavar e esterilizar o pote devolvido (estimativa).", "Planilha")

# ============================ Insumos ============================
ins = wb.create_sheet("Insumos")
titulo(ins, "Insumos", "O que vocês compram. Atualize preço e tamanho da compra; o custo por unidade é calculado.", 7)
cabecalho(ins, 4, ["Insumo", "Unidade", "Qtd na compra", "Preço pago", "Custo por unidade", "Origem do número", "Obs."],
          [30, 10, 13, 13, 15, 20, 40])
insumos = [
    ("Cachaça-base", "ml", 1000, 7.00, "Informado por você", ""),
    ("Mel", "g", 1000, 40.00, "Estimativa", "Pote de 1 kg"),
    ("Limão", "un", 10, 5.00, "Estimativa", ""),
    ("Alecrim do mato (ramo)", "un", 10, 5.00, "Estimativa", ""),
    ("Pote 40 ml com tampa", "un", 1, 2.00, "Informado por você", ""),
    ("Pote 240 ml com tampa", "un", 1, 4.40, "Pesquisa out/2026", "Com tampa metálica e lacre: R$ 3,92 (caixa c/ 48) a R$ 5,33 (kit c/ 12)"),
    ("Garrafa 1 L com tampa", "un", 1, 7.00, "Estimativa", ""),
    ("Garrafa 500 ml com tampa", "un", 1, 4.50, "Estimativa", "Produto planejado"),
    ("Rótulo pote pequeno", "un", 100, 30.00, "Estimativa", "Gráfica, cento"),
    ("Rótulo pote grande", "un", 100, 60.00, "Estimativa", "Gráfica, cento"),
    ("Rótulo + lacre da garrafa", "un", 1, 1.50, "Estimativa", ""),
    ("Caixa kraft (kit)", "un", 1, 4.00, "Estimativa", ""),
    ("Fita + tag (kit)", "un", 1, 1.00, "Estimativa", ""),
    ("Palha / papel seda (kit)", "un", 1, 0.50, "Estimativa", ""),
]
INS = {}
for i, (n, u, q, p, o, obs) in enumerate(insumos):
    r = 5 + i
    est = o == "Estimativa"
    put(ins, f"A{r}", n)
    put(ins, f"B{r}", u)
    put(ins, f"C{r}", q, INPUT, QTD, AMARELO if est else None)
    put(ins, f"D{r}", p, INPUT, RS, AMARELO if est else None)
    put(ins, f"E{r}", f"=IF(C{r}>0,D{r}/C{r},0)", FORM, '"R$" #,##0.0000')
    put(ins, f"F{r}", o)
    put(ins, f"G{r}", obs)
    INS[n] = f"Insumos!$E${r}"

# ============================ Produtos ============================
pr = wb.create_sheet("Produtos")
titulo(pr, "Produtos", "Custo unitário, preço, margem e desconto de refil de cada produto.", 17)
cabecalho(pr, 4, ["Código", "Produto", "Tamanho", "Status", "Conteúdo (ingredientes)", "Embalagem", "Rótulo / lacre", "Gás / energia",
                  "Custo unitário", "Preço sugerido", "Preço de venda", "Lucro por unidade", "Margem (%)", "Markup",
                  "Desconto refil (R$)", "Validade (dias)", "Obs."],
          [11, 34, 9, 11, 14, 12, 12, 12, 13, 13, 13, 13, 11, 10, 12, 11, 44])
sabores = [("G-PIM", "Geleia de Pimenta", 1.00), ("G-MAN", "Geleia de Manga com Maracujá", 1.00),
           ("PES", "Pesto de Manjericão", 1.80), ("CHU", "Chutney de Tamarindo", 1.20),
           ("MPI", "Molho de Pimenta", 0.90), ("ALH", "Alho Confit (em azeite)", 1.60), ("MOS", "Mostarda", 0.90)]
precos240 = {"PES": 32.90, "CHU": 32.90, "ALH": 32.90}
produtos = []
for cod, nome, c40 in sabores:
    produtos.append((f"{cod}-40", nome, "40 ml", "Ativo", c40, "Pote 40 ml com tampa", "Rótulo pote pequeno", 0.30, 11.90, 1.50))
for cod, nome, c40 in sabores:
    produtos.append((f"{cod}-240", nome, "240 ml", "Ativo", round(c40 * 6, 2), "Pote 240 ml com tampa", "Rótulo pote grande", 0.80,
                     precos240.get(cod, 29.90), 3.00))
produtos.append(("CAC-1000", "Cachaça com alecrim do mato, limão e mel", "1 L", "Ativo", "F1000", "Garrafa 1 L com tampa",
                 "Rótulo + lacre da garrafa", 0, 49.90, 5.00))
produtos.append(("CAC-500", "Cachaça com alecrim do mato, limão e mel", "500 ml", "Planejado", "F500", "Garrafa 500 ml com tampa",
                 "Rótulo + lacre da garrafa", 0, 32.90, 3.00))
P0 = 5
for i, (cod, nome, tam, st, cont, emb, rot, gas, preco, refil) in enumerate(produtos):
    r = P0 + i
    base = cod.split("-")[0]
    put(pr, f"A{r}", cod, fonte(bold=True))
    put(pr, f"B{r}", nome)
    put(pr, f"C{r}", tam)
    put(pr, f"D{r}", st, INPUT)
    if cont == "F1000":
        put(pr, f"E{r}", f"=1000*{INS['Cachaça-base']}+70*{INS['Mel']}+1*{INS['Limão']}+2*{INS['Alecrim do mato (ramo)']}", LINK, RS)
        pr[f"E{r}"].comment = Comment("1 L de cachaça-base + 70 g de mel + 1 limão + 2 ramos de alecrim (receita estimada: ajuste as quantidades na fórmula).", "Planilha")
    elif cont == "F500":
        put(pr, f"E{r}", f"=E{r - 1}/2", FORM, RS)
    else:
        put(pr, f"E{r}", cont, INPUT, RS, AMARELO)
    put(pr, f"F{r}", f"={INS[emb]}", LINK, RS)
    put(pr, f"G{r}", f"={INS[rot]}", LINK, RS)
    put(pr, f"H{r}", gas, INPUT, RS, AMARELO if gas else None)
    put(pr, f"I{r}", f"=SUM(E{r}:H{r})", FORM, RS)
    put(pr, f"J{r}", f"=IF(I{r}>0,CEILING(I{r}*Parâmetros!$F$6,1)-0.1,0)", FORM, RS)
    put(pr, f"K{r}", preco, INPUT, RS, AMARELO)
    put(pr, f"L{r}", f"=K{r}-I{r}", FORM, RS)
    put(pr, f"M{r}", f"=IF(K{r}>0,L{r}/K{r},0)", FORM, PCT)
    put(pr, f"N{r}", f"=IF(I{r}>0,K{r}/I{r},0)", FORM, X)
    put(pr, f"O{r}", refil, INPUT, RS, AMARELO)
    put(pr, f"P{r}", 15 if base in ("PES", "ALH") else 180, INPUT, QTD, AMARELO)
    obs = ""
    if base in ("PES", "ALH"):
        obs = "Conservar refrigerado (alimento em óleo)"
    if base == "CAC":
        obs = "Só para maiores de 18 anos. Verificar registro no MAPA (bebida mista)."
    put(pr, f"Q{r}", obs)
P1 = P0 + len(produtos) - 1
NPROD = len(produtos)
pr[f"E{P0}"].comment = Comment("Estimativa do custo dos ingredientes de 1 pote. Calcule: custo da panelada ÷ número de potes que ela rende.", "Planilha")
pr[f"O{P0}"].comment = Comment("Desconto dado quando o cliente devolve o pote para encher de novo. Deixe menor que o custo do pote, para o refil sair mais barato para os dois.", "Planilha")
pr.conditional_formatting.add(f"M{P0}:M{P1}", CellIsRule(operator="lessThan", formula=["0.5"], font=Font(name=F, color="C00000", bold=True)))
pr.conditional_formatting.add(f"A{P0}:Q{P1}", FormulaRule(formula=[f'$D{P0}="Planejado"'], fill=SUAVE))
put(pr, f"A{P1 + 2}", "Margem abaixo de 50% fica em vermelho. Linhas em bege = produtos planejados (garrafa 500 ml).",
    fonte(9, italic=True), borda=False)
dv_st = DataValidation(type="list", formula1='"Ativo,Planejado,Pausado"', allow_blank=True)
pr.add_data_validation(dv_st)
dv_st.add(f"D{P0}:D{P1}")

# ============================ Kits ============================
kt = wb.create_sheet("Kits")
c0, cN = 4, 4 + NPROD - 1
titulo(kt, "Kits", "Composição (quantos de cada produto), custo, preço e margem de cada kit.", cN + 5)
kt["A3"] = "Custo unitário de cada produto →"
kt["A3"].font = fonte(9, italic=True)
nomes = ["Código", "Kit", "Embalagem do kit"] + [p[0] for p in produtos] + ["Custo do kit", "Preço de venda", "Lucro", "Margem (%)", "Vendidos"]
cabecalho(kt, 5, nomes, [10, 34, 13] + [9] * NPROD + [13, 13, 12, 11, 10])
for j in range(NPROD):
    col = L(c0 + j)
    put(kt, f"{col}4", f"=INDEX(Produtos!$I${P0}:$I${P1},MATCH({col}5,Produtos!$A${P0}:$A${P1},0))", LINK, '"R$" #,##0.00')
kits = [
    ("K-DEG", "Degustação (4 minis 40 ml)", ["Caixa kraft (kit)", "Fita + tag (kit)"], {"G-PIM-40": 1, "G-MAN-40": 1, "CHU-40": 1, "MOS-40": 1}, 39.90),
    ("K-TAB", "Tábua de petiscos", ["Caixa kraft (kit)", "Fita + tag (kit)", "Palha / papel seda (kit)"],
     {"G-PIM-240": 1, "CHU-240": 1, "ALH-40": 1}, 74.90),
    ("K-COMP", "Receitas do Bistrô completo (7 minis + 1 grande)", ["Caixa kraft (kit)", "Fita + tag (kit)", "Palha / papel seda (kit)"],
     {**{f"{s[0]}-40": 1 for s in sabores}, "G-PIM-240": 1}, 104.90),
    ("K-CAC", "Presente com cachaça (500 ml + 2 grandes)", ["Caixa kraft (kit)", "Fita + tag (kit)", "Palha / papel seda (kit)"],
     {"CAC-500": 1, "G-PIM-240": 1, "CHU-240": 1}, 94.90),
    ("K-LEMB", "Lembrancinha (1 mini com fita e tag)", ["Fita + tag (kit)"], {"G-PIM-40": 1}, 9.90),
]
K0 = 6
cc, cp, cl, cm, cv = (L(cN + k) for k in range(1, 6))
for i, (cod, nome, emb, comp, preco) in enumerate(kits):
    r = K0 + i
    put(kt, f"A{r}", cod, fonte(bold=True))
    put(kt, f"B{r}", nome)
    put(kt, f"C{r}", "=" + "+".join(INS[e] for e in emb), LINK, RS)
    for j, p in enumerate(produtos):
        put(kt, f"{L(c0 + j)}{r}", comp.get(p[0], 0), INPUT, QTD)
    put(kt, f"{cc}{r}", f"=C{r}+SUMPRODUCT({L(c0)}{r}:{L(cN)}{r},{L(c0)}$4:{L(cN)}$4)", FORM, RS)
    put(kt, f"{cp}{r}", preco, INPUT, RS, AMARELO)
    put(kt, f"{cl}{r}", f"={cp}{r}-{cc}{r}", FORM, RS)
    put(kt, f"{cm}{r}", f"=IF({cp}{r}>0,{cl}{r}/{cp}{r},0)", FORM, PCT)
    put(kt, f"{cv}{r}", f"=SUMIFS(Vendas!$D$5:$D$604,Vendas!$B$5:$B$604,A{r})", FORM, QTD)
K1 = K0 + len(kits) - 1
put(kt, f"A{K1 + 2}", "O kit com cachaça usa a garrafa de 500 ml (planejada). Mude a composição à vontade (0, 1, 2…).",
    fonte(9, italic=True), borda=False)
# quantos de cada produto já saíram dentro de kits vendidos (usado pela aba Estoque)
KUSO = K1 + 4
put(kt, f"B{KUSO}", "Saíram em kits vendidos →", fonte(bold=True))
for j in range(NPROD):
    col = L(c0 + j)
    put(kt, f"{col}{KUSO}", f"=SUMPRODUCT({col}{K0}:{col}{K1},${cv}${K0}:${cv}${K1})", FORM, QTD)
kt.freeze_panes = "D6"

# ============================ Catálogo (Parâmetros H:M) ============================
for col, n, w in [("H", "Catálogo (código)", 16), ("I", "Item", 46), ("J", "Preço", 11), ("K", "Custo", 11),
                  ("L", "Embalagem", 11), ("M", "Desconto refil", 12)]:
    put(pm, f"{col}4", n, fonte(bold=True, color=CREME), fill=TOPO)
    pm[f"{col}4"].alignment = Alignment(horizontal="center", wrap_text=True)
    pm.column_dimensions[col].width = w
pm.column_dimensions["G"].width = 3
k = 5
for i in range(NPROD):
    r = P0 + i
    put(pm, f"H{k}", f"=Produtos!A{r}", LINK)
    put(pm, f"I{k}", f'=Produtos!B{r}&" "&Produtos!C{r}', LINK)
    put(pm, f"J{k}", f"=Produtos!K{r}", LINK, RS)
    put(pm, f"K{k}", f"=Produtos!I{r}", LINK, RS)
    put(pm, f"L{k}", f"=Produtos!F{r}", LINK, RS)
    put(pm, f"M{k}", f"=Produtos!O{r}", LINK, RS)
    k += 1
for i in range(len(kits)):
    r = K0 + i
    put(pm, f"H{k}", f"=Kits!A{r}", LINK)
    put(pm, f"I{k}", f'="Kit "&Kits!B{r}', LINK)
    put(pm, f"J{k}", f"=Kits!{cp}{r}", LINK, RS)
    put(pm, f"K{k}", f"=Kits!{cc}{r}", LINK, RS)
    put(pm, f"L{k}", 0, FORM, RS)
    put(pm, f"M{k}", 0, FORM, RS)
    k += 1
CAT_FIM = k - 1
CAT = f"Parâmetros!$H$5:$H${CAT_FIM}"
CATI, CATP, CATC, CATE, CATR = (f"Parâmetros!${c}$5:${c}${CAT_FIM}" for c in "IJKLM")


def busca(col, cod):
    return f"IFERROR(INDEX({col},MATCH({cod},{CAT},0)),0)"


# ============================ Vendas ============================
vd = wb.create_sheet("Vendas")
titulo(vd, "Vendas", 'Uma linha por venda. Preencha só as colunas azuis. Brinde: canal "Brinde". Refil: "Sim" quando o cliente traz o pote.', 17)
cabecalho(vd, 4, ["Data", "Código", "Item", "Qtd", "Preço de tabela", "Desconto (R$)", "Pagamento", "Canal", "Refil?",
                  "Desconto refil", "Bruto", "Taxa", "Líquido", "Custo", "Lucro", "Ano", "Mês"],
          [12, 11, 40, 6, 12, 11, 17, 17, 8, 11, 12, 10, 12, 11, 12, 7, 6])
NV = 600
for r in range(5, 5 + NV):
    for c in "ABDFGHI":
        vd[f"{c}{r}"].font = INPUT
        vd[f"{c}{r}"].border = BORDA
    vd[f"A{r}"].number_format = DATA
    vd[f"D{r}"].number_format = QTD
    vd[f"F{r}"].number_format = RS
    vd[f"I{r}"].alignment = Alignment(horizontal="center")
    put(vd, f"C{r}", f'=IF(B{r}="","",IFERROR(INDEX({CATI},MATCH(B{r},{CAT},0)),"código não encontrado"))')
    put(vd, f"E{r}", f'=IF(B{r}="","",{busca(CATP, f"B{r}")})', FORM, RS)
    put(vd, f"J{r}", f'=IF(OR(B{r}="",I{r}<>"Sim"),0,D{r}*{busca(CATR, f"B{r}")})', FORM, RS)
    put(vd, f"K{r}", f'=IF(B{r}="","",IF(H{r}="Brinde",0,MAX(0,D{r}*E{r}-F{r}-J{r})))', FORM, RS)
    put(vd, f"L{r}", f'=IF(B{r}="","",K{r}*IFERROR(INDEX(Parâmetros!$B$5:$B$9,MATCH(G{r},Parâmetros!$A$5:$A$9,0)),0))', FORM, RS)
    put(vd, f"M{r}", f'=IF(B{r}="","",K{r}-L{r})', FORM, RS)
    # refil: o pote volta (sai a embalagem do custo) e entra a higienização
    put(vd, f"N{r}", f'=IF(B{r}="","",D{r}*({busca(CATC, f"B{r}")}-IF(I{r}="Sim",{busca(CATE, f"B{r}")}-Parâmetros!$F$12,0)))', FORM, RS)
    put(vd, f"O{r}", f'=IF(B{r}="","",M{r}-N{r})', FORM, RS)
    put(vd, f"P{r}", f'=IF(A{r}="","",YEAR(A{r}))', FORM, "0")
    put(vd, f"Q{r}", f'=IF(A{r}="","",MONTH(A{r}))', FORM, "0")
VU = 4 + NV
for formula, rng in [(f"={CAT}", f"B5:B{VU}"), ("=Parâmetros!$A$5:$A$9", f"G5:G{VU}"),
                     ("=Parâmetros!$D$5:$D$9", f"H5:H{VU}"), ('"Sim,Não"', f"I5:I{VU}")]:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    vd.add_data_validation(dv)
    dv.add(rng)
vd["A5"] = dt.date(2026, 10, 22)
vd["B5"] = "CAC-1000"
vd["D5"] = 1
vd["F5"] = 0
vd["G5"] = "Pix"
vd["H5"] = "Direct Instagram"
vd["I5"] = "Não"
vd["A5"].comment = Comment("Linha de EXEMPLO: apague ou escreva por cima.", "Planilha")

# ============================ Movimentos ============================
mv = wb.create_sheet("Movimentos")
titulo(mv, "Movimentos de estoque", "Produção (+), perdas e ajustes (+ ou −). Vendas e brindes NÃO entram aqui: já baixam o estoque pela aba Vendas.", 8)
cabecalho(mv, 4, ["Data", "Código", "Item", "Tipo", "Quantidade (+/−)", "Lote", "Validade", "Obs."], [12, 12, 42, 12, 14, 12, 12, 36])
PRODC = f"Produtos!$A${P0}:$A${P1}"
for r in range(5, 405):
    for c in "ABDEFH":
        mv[f"{c}{r}"].font = INPUT
        mv[f"{c}{r}"].border = BORDA
    mv[f"A{r}"].number_format = DATA
    mv[f"E{r}"].number_format = QTD
    put(mv, f"C{r}", f'=IF(B{r}="","",IFERROR(INDEX({CATI},MATCH(B{r},{CAT},0)),"código não encontrado"))')
    put(mv, f"G{r}", f'=IF(OR(A{r}="",B{r}="",D{r}<>"Produção"),"",A{r}+IFERROR(INDEX(Produtos!$P${P0}:$P${P1},MATCH(B{r},{PRODC},0)),0))', FORM, DATA)
for formula, rng in [('"Produção,Perda,Ajuste"', "D5:D404"), (f"={PRODC}", "B5:B404")]:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    mv.add_data_validation(dv)
    dv.add(rng)
mv["A5"] = dt.date(2026, 10, 20)
mv["B5"] = "CAC-1000"
mv["D5"] = "Produção"
mv["E5"] = 10
mv["F5"] = "L001"
mv["H5"] = "Exemplo: apague ou escreva por cima"
mv["G4"].comment = Comment("Calculada: data de produção + validade (dias) da aba Produtos.", "Planilha")

# ============================ Estoque ============================
es = wb.create_sheet("Estoque")
titulo(es, "Estoque", "Saldo de cada produto. Só o estoque inicial e o mínimo são digitados.", 12)
cabecalho(es, 4, ["Código", "Item", "Estoque inicial", "Produzido", "Perdas / ajustes", "Vendido avulso (inclui brindes e refis)",
                  "Saiu em kits", "Estoque atual", "Mínimo", "Situação", "Custo unitário", "Valor em estoque"],
          [11, 44, 11, 11, 12, 16, 11, 11, 9, 11, 12, 14])
KCOMP = f"Kits!${L(c0)}${K0}:${L(cN)}${K1}"
KCAB = f"Kits!${L(c0)}$5:${L(cN)}$5"
KVEND = f"Kits!${cv}${K0}:${cv}${K1}"
for i in range(NPROD):
    r, p = 5 + i, P0 + i
    put(es, f"A{r}", f"=Produtos!A{p}", LINK)
    put(es, f"B{r}", f'=Produtos!B{p}&" "&Produtos!C{p}', LINK)
    put(es, f"C{r}", 0, INPUT, QTD)
    put(es, f"D{r}", f'=SUMIFS(Movimentos!$E$5:$E$404,Movimentos!$B$5:$B$404,A{r},Movimentos!$D$5:$D$404,"Produção")', FORM, QTD)
    put(es, f"E{r}", f'=SUMIFS(Movimentos!$E$5:$E$404,Movimentos!$B$5:$B$404,A{r},Movimentos!$D$5:$D$404,"<>Produção")', FORM, QTD)
    put(es, f"F{r}", f"=SUMIFS(Vendas!$D$5:$D$604,Vendas!$B$5:$B$604,A{r})", FORM, QTD)
    put(es, f"G{r}", f"=IFERROR(INDEX(Kits!${L(c0)}${KUSO}:${L(cN)}${KUSO},MATCH(A{r},{KCAB},0)),0)", LINK, QTD)
    put(es, f"H{r}", f"=C{r}+D{r}+E{r}-F{r}-G{r}", fonte(bold=True), QTD)
    put(es, f"I{r}", 5 if produtos[i][3] == "Ativo" else 0, INPUT, QTD, AMARELO)
    put(es, f"J{r}", f'=IF(Produtos!D{p}<>"Ativo","—",IF(H{r}<0,"Conferir",IF(H{r}<=I{r},"Repor","OK")))')
    put(es, f"K{r}", f"=Produtos!I{p}", LINK, RS)
    put(es, f"L{r}", f"=MAX(0,H{r})*K{r}", FORM, RS)
EU = 4 + NPROD
put(es, f"K{EU + 1}", "Total", fonte(bold=True))
put(es, f"L{EU + 1}", f"=SUM(L5:L{EU})", fonte(bold=True), RS)
put(es, f"A{EU + 3}", "Refil: o pote volta, mas o produto sai do estoque normalmente (a geleia foi feita e envasada do mesmo jeito).",
    fonte(9, italic=True), borda=False)
es.conditional_formatting.add(f"J5:J{EU}", CellIsRule(operator="equal", formula=['"Repor"'], fill=PatternFill("solid", fgColor="F8C9C2"),
                                                       font=Font(name=F, bold=True, color=VINHO)))
es.conditional_formatting.add(f"J5:J{EU}", CellIsRule(operator="equal", formula=['"OK"'], font=Font(name=F, color="396600")))
es.conditional_formatting.add(f"J5:J{EU}", CellIsRule(operator="equal", formula=['"Conferir"'], fill=AMARELO))

# ============================ Despesas ============================
dp = wb.create_sheet("Despesas")
titulo(dp, "Despesas", "Gastos que não são custo de produto (os insumos já estão no custo de cada pote ou garrafa).", 6)
cabecalho(dp, 4, ["Data", "Descrição", "Categoria", "Valor", "Ano", "Mês"], [12, 40, 20, 12, 7, 6])
for r in range(5, 305):
    for c in "ABCD":
        dp[f"{c}{r}"].font = INPUT
        dp[f"{c}{r}"].border = BORDA
    dp[f"A{r}"].number_format = DATA
    dp[f"D{r}"].number_format = RS
    put(dp, f"E{r}", f'=IF(A{r}="","",YEAR(A{r}))', FORM, "0")
    put(dp, f"F{r}", f'=IF(A{r}="","",MONTH(A{r}))', FORM, "0")
dv = DataValidation(type="list", formula1='"Divulgação,Gráfica,Alvará / taxas,Entrega,Equipamentos,Outros"', allow_blank=True)
dp.add_data_validation(dv)
dv.add("C5:C304")

# ============================ Resumo ============================
rs = wb.create_sheet("Resumo")
titulo(rs, "Resumo", "Tudo calculado a partir de Vendas e Despesas. Para ver outros meses, mude o ano e o mês (em azul).", 10)
put(rs, "A4", "Por mês", fonte(12, bold=True, color=VINHO), borda=False)
cab = ["Ano", "Mês", "Bruto", "Descontos de refil", "Taxas", "Líquido", "Custo dos produtos", "Lucro das vendas", "Despesas", "Resultado do mês"]
for i, (n, w) in enumerate(zip(cab, [8, 44, 13, 12, 11, 13, 14, 14, 12, 15]), 1):
    c = rs.cell(5, i, n)
    c.font = fonte(bold=True, color=CREME)
    c.fill = TOPO
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDA
    rs.column_dimensions[L(i)].width = w
rs.row_dimensions[5].height = 30
ano, mes = 2026, 10
for i in range(15):
    r = 6 + i
    put(rs, f"A{r}", ano, INPUT, "0")
    put(rs, f"B{r}", mes, INPUT, "0")
    rs[f"B{r}"].alignment = Alignment(horizontal="center")
    crit = f"Vendas!$P$5:$P$604,$A{r},Vendas!$Q$5:$Q$604,$B{r}"
    put(rs, f"C{r}", f"=SUMIFS(Vendas!$K$5:$K$604,{crit})", FORM, RS)
    put(rs, f"D{r}", f"=SUMIFS(Vendas!$J$5:$J$604,{crit})", FORM, RS)
    put(rs, f"E{r}", f"=SUMIFS(Vendas!$L$5:$L$604,{crit})", FORM, RS)
    put(rs, f"F{r}", f"=SUMIFS(Vendas!$M$5:$M$604,{crit})", FORM, RS)
    put(rs, f"G{r}", f"=SUMIFS(Vendas!$N$5:$N$604,{crit})", FORM, RS)
    put(rs, f"H{r}", f"=F{r}-G{r}", FORM, RS)
    put(rs, f"I{r}", f"=SUMIFS(Despesas!$D$5:$D$304,Despesas!$E$5:$E$304,$A{r},Despesas!$F$5:$F$304,$B{r})", FORM, RS)
    put(rs, f"J{r}", f"=H{r}-I{r}", fonte(bold=True), RS)
    mes += 1
    if mes > 12:
        mes, ano = 1, ano + 1
t = 21
put(rs, f"A{t}", "Total", fonte(bold=True))
put(rs, f"B{t}", "")
for c in "CDEFGHIJ":
    put(rs, f"{c}{t}", f"=SUM({c}6:{c}{t - 1})", fonte(bold=True), RS)
rs.conditional_formatting.add(f"J6:J{t}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="C00000")))

p0 = t + 3
put(rs, f"A{p0}", "Por produto e kit", fonte(12, bold=True, color=VINHO), borda=False)
cab2 = ["Código", "Item", "Qtd vendida", "Bruto", "Líquido", "Custo", "Lucro", "Margem (%)", "Brindes dados", "Refis"]
for i, n in enumerate(cab2, 1):
    c = rs.cell(p0 + 1, i, n)
    c.font = fonte(bold=True, color=CREME)
    c.fill = TOPO
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    c.border = BORDA
for i in range(CAT_FIM - 4):
    r = p0 + 2 + i
    put(rs, f"A{r}", f"=Parâmetros!H{5 + i}", LINK)
    put(rs, f"B{r}", f"=Parâmetros!I{5 + i}", LINK)
    put(rs, f"C{r}", f"=SUMIFS(Vendas!$D$5:$D$604,Vendas!$B$5:$B$604,$A{r})", FORM, QTD)
    put(rs, f"D{r}", f"=SUMIFS(Vendas!$K$5:$K$604,Vendas!$B$5:$B$604,$A{r})", FORM, RS)
    put(rs, f"E{r}", f"=SUMIFS(Vendas!$M$5:$M$604,Vendas!$B$5:$B$604,$A{r})", FORM, RS)
    put(rs, f"F{r}", f"=SUMIFS(Vendas!$N$5:$N$604,Vendas!$B$5:$B$604,$A{r})", FORM, RS)
    put(rs, f"G{r}", f"=E{r}-F{r}", FORM, RS)
    put(rs, f"H{r}", f"=IF(D{r}>0,G{r}/D{r},0)", FORM, PCT)
    put(rs, f"I{r}", f'=SUMIFS(Vendas!$D$5:$D$604,Vendas!$B$5:$B$604,$A{r},Vendas!$H$5:$H$604,"Brinde")', FORM, QTD)
    put(rs, f"J{r}", f'=SUMIFS(Vendas!$D$5:$D$604,Vendas!$B$5:$B$604,$A{r},Vendas!$I$5:$I$604,"Sim")', FORM, QTD)
rs.freeze_panes = "A6"

ordem = ["Leia-me", "Vendas", "Movimentos", "Estoque", "Resumo", "Produtos", "Kits", "Insumos", "Despesas", "Parâmetros"]
wb._sheets = [wb[n] for n in ordem]
cor = {"Leia-me": VINHO, "Vendas": "396600", "Movimentos": "396600", "Despesas": "396600", "Estoque": "003984", "Resumo": "003984"}
for w in wb.worksheets:
    w.sheet_properties.tabColor = cor.get(w.title, "D16F4A")
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(SAIDA)
print("ok", NPROD, "produtos,", len(kits), "kits")
