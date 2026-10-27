"""Aba "Projeção" da planilha de controle das Receitas do Bistrô: investimento inicial, 3 cenários de venda,
12 meses de resultado acumulado, payback e gráfico. Lê preço e custo das abas Produtos, Kits e Insumos.

    python receitas/projecao_planilha.py            # acrescenta (ou refaz) a aba no arquivo existente, sem tocar nas outras
"""
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ARQ = r"C:/LIA - UEFS/motores-video/saida/receitas-bistro/Receitas-do-Bistro-controle.xlsx"
F = "Arial"
VINHO, CREME, MARROM = "8D0C00", "FFFDEA", "3D2F29"
INPUT, LINK, FORM = Font(name=F, size=10, color="0000FF"), Font(name=F, size=10, color="008000"), Font(name=F, size=10)
NEG = Font(name=F, size=10, bold=True)
AMARELO = PatternFill("solid", fgColor="FFF2A8")
TOPO = PatternFill("solid", fgColor=MARROM)
fino = Side(style="thin", color="D9CFB8")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
RS = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
RS0 = '"R$" #,##0;-"R$" #,##0;"-"'
PCT = '0%;-0%;"-"'
QTD = '#,##0;-#,##0;"-"'
# cenários: cor validada (dataviz: lightness, croma, daltonismo e contraste) + traço diferente em cada linha
CENARIOS = [("Pessimista", "E", "C4552B", "dash"), ("Realista", "F", "2F6DB5", "solid"), ("Otimista", "G", "4F8A1E", "sysDot")]

PROD = "Produtos!$A$5:$A$20"
PC, PS, PK = "Produtos!$C$5:$C$20", "Produtos!$I$5:$I$20", "Produtos!$K$5:$K$20"
INS_N, INS_C = "Insumos!$A$5:$A$18", "Insumos!$E$5:$E$18"


def put(ws, ref, v, font=FORM, fmt=None, fill=None, borda=True):
    c = ws[ref]
    c.value, c.font = v, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if borda:
        c.border = BORDA
    return c


def cab(ws, linha, nomes):
    for i, n in enumerate(nomes):
        c = ws.cell(linha, 1 + i, n)
        c.font, c.fill, c.border = Font(name=F, size=10, bold=True, color=CREME), TOPO, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 32


def secao(ws, ref, texto):
    put(ws, ref, texto, Font(name=F, size=12, bold=True, color=VINHO), borda=False)


def montar(wb):
    if "Projeção" in wb.sheetnames:
        del wb["Projeção"]
    ws = wb.create_sheet("Projeção", index=wb.sheetnames.index("Resumo") + 1)
    ws.sheet_properties.tabColor = "003984"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Projeção · quanto investir e quanto pode voltar"
    ws["A1"].font = Font(name=F, size=14, bold=True, color=CREME)
    for c in range(1, 11):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=VINHO)
    ws["A2"] = "Estimativa, não promessa: mude as quantidades em azul e veja os 3 cenários. Preço e custo vêm das abas Produtos e Kits."
    ws["A2"].font = Font(name=F, size=9, italic=True, color=MARROM)
    for col, w in zip("ABCDEFGHIJ", (44, 13, 13, 14, 13, 13, 13, 13, 13, 13)):
        ws.column_dimensions[col].width = w

    # 1. investimento
    secao(ws, "A4", "1. Investimento inicial")
    cab(ws, 5, ["Item", "Quantidade", "Custo unitário", "Total", "Obs."])
    ws.merge_cells("E5:J5")
    itens = [
        ("Primeira produção · minis 40 ml", 70, f'=AVERAGEIFS({PS},{PC},"40 ml")', LINK, "10 de cada sabor (custo médio da aba Produtos)"),
        ("Primeira produção · potes 240 ml", 35, f'=AVERAGEIFS({PS},{PC},"240 ml")', LINK, "5 de cada sabor"),
        ("Primeira produção · cachaça 1 L", 10, f'=INDEX({PS},MATCH("CAC-1000",{PROD},0))', LINK, ""),
        ("Caixas, fitas e tags para kits", 10,
         f'=INDEX({INS_C},MATCH("Caixa kraft (kit)",{INS_N},0))+INDEX({INS_C},MATCH("Fita + tag (kit)",{INS_N},0))', LINK, ""),
        ("Utensílios (funil, concha, balança, etiquetadora)", 1, 300, INPUT, "ESTIMATIVA: troque pelo que precisar comprar"),
        ("Divulgação de lançamento (impulsionar posts)", 1, 150, INPUT, "ESTIMATIVA"),
        ("Legalização (alvará, rótulo nutricional, registro da cachaça)", 1, 0, INPUT, "PREENCHER: consulte a vigilância e o MAPA"),
    ]
    for i, (n, q, cu, fcu, obs) in enumerate(itens):
        r = 6 + i
        put(ws, f"A{r}", n)
        put(ws, f"B{r}", q, INPUT, QTD, AMARELO)
        put(ws, f"C{r}", cu, fcu, RS, AMARELO if fcu is INPUT else None)
        put(ws, f"D{r}", f"=B{r}*C{r}", FORM, RS)
        put(ws, f"E{r}", obs, Font(name=F, size=9, italic=True), borda=False)
    put(ws, "A13", "Total do investimento", NEG)
    put(ws, "D13", "=SUM(D6:D12)", NEG, RS)
    INV = "$D$13"

    # 2. premissas
    secao(ws, "A15", "2. Premissas de venda")
    cab(ws, 16, ["Linha", "Preço médio", "Custo médio", "Ganho por unidade (após taxa)", "Pessimista", "Realista", "Otimista"])
    ws["E16"].comment = Comment("Unidades vendidas no 1º mês em cada cenário. A partir daí crescem pela taxa da linha 21.", "Planilha")
    linhas = [
        ("Minis 40 ml", f'=AVERAGEIFS({PK},{PC},"40 ml")', f'=AVERAGEIFS({PS},{PC},"40 ml")', (20, 40, 70)),
        ("Potes 240 ml", f'=AVERAGEIFS({PK},{PC},"240 ml")', f'=AVERAGEIFS({PS},{PC},"240 ml")', (10, 25, 45)),
        ("Cachaça 1 L", f'=INDEX({PK},MATCH("CAC-1000",{PROD},0))', f'=INDEX({PS},MATCH("CAC-1000",{PROD},0))', (4, 10, 20)),
        ("Kits (média dos kits)", "=AVERAGE(Kits!$U$6:$U$10)", "=AVERAGE(Kits!$T$6:$T$10)", (2, 6, 15)),
    ]
    for i, (n, p, c, vols) in enumerate(linhas):
        r = 17 + i
        put(ws, f"A{r}", n)
        put(ws, f"B{r}", p, LINK, RS)
        put(ws, f"C{r}", c, LINK, RS)
        put(ws, f"D{r}", f"=B{r}*(1-$B$23)-C{r}", FORM, RS)
        for j, v in enumerate(vols):
            put(ws, f"{'EFG'[j]}{r}", v, INPUT, QTD, AMARELO)
    put(ws, "A21", "Crescimento das vendas por mês")
    put(ws, "A22", "Despesas fixas por mês (divulgação, entregas, gráfica…)")
    for j, (g, d) in enumerate([(0, 150), (0.05, 150), (0.10, 200)]):
        put(ws, f"{'EFG'[j]}21", g, INPUT, PCT, AMARELO)
        put(ws, f"{'EFG'[j]}22", d, INPUT, RS0, AMARELO)
    put(ws, "A23", "Taxa média das formas de pagamento")
    put(ws, "B23", 0.02, INPUT, '0.0%', AMARELO)
    ws["B23"].comment = Comment("Mistura de Pix (0%) e cartão (2 a 5%). Ajuste conforme o que os clientes mais usam.", "Planilha")

    # 3. mês a mês
    secao(ws, "A25", "3. Mês a mês (12 meses)")
    nomes = ["Mês"]
    for n, *_ in CENARIOS:
        nomes += [f"{n}: faturamento", f"{n}: resultado", f"{n}: saldo acumulado"]
    cab(ws, 26, nomes)
    for m in range(1, 13):
        r = 26 + m
        put(ws, f"A{r}", m, FORM, '"Mês "0')
        for j, (n, vcol, *_) in enumerate(CENARIOS):
            cf, cr, ca = (chr(ord("B") + 3 * j + k) for k in range(3))
            cresc = f"(1+${vcol}$21)^($A{r}-1)"
            put(ws, f"{cf}{r}", f"=SUMPRODUCT($B$17:$B$20,${vcol}$17:${vcol}$20)*{cresc}", FORM, RS0)
            put(ws, f"{cr}{r}", f"=SUMPRODUCT($D$17:$D$20,${vcol}$17:${vcol}$20)*{cresc}-${vcol}$22", FORM, RS0)
            put(ws, f"{ca}{r}", f"=-{INV}+{cr}{r}" if m == 1 else f"={ca}{r - 1}+{cr}{r}", NEG, RS0)
    ws.conditional_formatting.add("B27:J38", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="C00000")))

    # 4. resumo dos cenários
    secao(ws, "A41", "4. Resumo dos cenários")
    cab(ws, 42, ["Cenário", "Faturamento em 12 meses", "Lucro em 12 meses", "Investimento", "Saldo ao fim de 12 meses",
                 "Retorno sobre o investimento", "O investimento se paga no"])
    for j, (n, vcol, *_) in enumerate(CENARIOS):
        r = 43 + j
        cf, cr, ca = (chr(ord("B") + 3 * j + k) for k in range(3))
        put(ws, f"A{r}", n, NEG)
        put(ws, f"B{r}", f"=SUM({cf}27:{cf}38)", FORM, RS0)
        put(ws, f"C{r}", f"=SUM({cr}27:{cr}38)", FORM, RS0)
        put(ws, f"D{r}", f"={INV}", FORM, RS0)
        put(ws, f"E{r}", f"={ca}38", NEG, RS0)
        put(ws, f"F{r}", f"=IF({INV}>0,E{r}/{INV},0)", FORM, PCT)
        put(ws, f"G{r}", f'=IF(COUNTIF({ca}27:{ca}38,"<0")>=12,"depois de 12 meses","mês "&(COUNTIF({ca}27:{ca}38,"<0")+1))', NEG)
        ws[f"G{r}"].alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("E43:E45", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="C00000")))
    put(ws, "A47", "Para pagar as despesas fixas do cenário realista, é preciso vender por mês o equivalente a", borda=False)
    put(ws, "A48", "  minis de 40 ml", borda=False)
    put(ws, "B48", "=IF($D$17>0,ROUNDUP($F$22/$D$17,0),0)", NEG, QTD)
    put(ws, "A49", "  ou potes de 240 ml", borda=False)
    put(ws, "B49", "=IF($D$18>0,ROUNDUP($F$22/$D$18,0),0)", NEG, QTD)
    put(ws, "A50", "  ou garrafas de cachaça", borda=False)
    put(ws, "B50", "=IF($D$19>0,ROUNDUP($F$22/$D$19,0),0)", NEG, QTD)
    put(ws, "A52", "Saldo acumulado começa negativo (o investimento) e sobe a cada mês com o resultado. Quando cruza o zero, o investimento se pagou.",
        Font(name=F, size=9, italic=True), borda=False)

    # gráfico: saldo acumulado por cenário (um eixo, legenda, traço diferente em cada linha)
    ch = LineChart()
    ch.title = "Saldo acumulado (investimento + resultado de cada mês)"
    ch.y_axis.title = "R$"
    ch.y_axis.numFmt = '"R$" #,##0'
    ch.y_axis.majorGridlines.spPr = None
    ch.x_axis.title = "Mês"
    ch.height, ch.width = 9, 18
    ch.legend.position = "b"
    for j, (n, _, cor, traco) in enumerate(CENARIOS):
        col = 4 + 3 * j
        ch.add_data(Reference(ws, min_col=col, min_row=27, max_row=38), titles_from_data=False)
        s = ch.series[-1]
        s.tx = SeriesLabel(v=n)
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = 28575
        s.graphicalProperties.line.dashStyle = traco
        s.smooth = False
    ch.set_categories(Reference(ws, min_col=1, min_row=27, max_row=38))
    ws.add_chart(ch, "L4")
    ws.freeze_panes = "B4"
    return ws


if __name__ == "__main__":
    wb = load_workbook(ARQ)
    montar(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(ARQ)
    print("aba Projeção pronta")
